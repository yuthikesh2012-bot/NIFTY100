"""
src/reports/export_peer_comparison.py

Day 20 — Generates output/peer_comparison.xlsx: 11 sheets (one per peer
group), company_id + company_name + 20 metric columns + percentile rank
columns for the 10 ranked metrics, colour-coded by percentile band,
benchmark company row highlighted, summary row with peer-group medians.
"""

from __future__ import annotations
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.analytics.peer import PEER_METRICS
from src.reports.export_screener import DISPLAY_COLUMNS, HEADER_FILL, HEADER_FONT, BODY_FONT

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GOLD = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
SUMMARY_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

# Raw KPI columns shown (excludes identifiers already handled separately).
METRIC_COLUMNS = [c for c, _ in DISPLAY_COLUMNS if c not in ("company_id", "company_name")]
METRIC_LABELS = {c: label for c, label in DISPLAY_COLUMNS}


def _percentile_fill(pct: float) -> PatternFill | None:
    if pct is None or pd.isna(pct):
        return None
    if pct >= 0.75:
        return GREEN
    if pct <= 0.25:
        return RED
    return YELLOW


def write_peer_group_sheet(wb: Workbook, group_name: str, group_df: pd.DataFrame,
                            percentiles_df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title=group_name[:31])

    pct_pivot = percentiles_df[percentiles_df["peer_group_name"] == group_name].pivot_table(
        index="company_id", columns="metric", values="percentile_rank", aggfunc="first"
    )

    metric_names = list(PEER_METRICS.keys())
    headers = ["Company", "Company Name"] + [METRIC_LABELS.get(c, c) for c in METRIC_COLUMNS] + \
              [f"{m} %ile" for m in metric_names]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    numeric_metric_cols = [c for c in METRIC_COLUMNS if c not in ("company_id", "company_name",
                                                                    "broad_sector", "year", "icr_label")]

    row_idx = 2
    for _, row in group_df.sort_values("company_id").iterrows():
        cid = row["company_id"]
        is_benchmark = bool(row.get("is_benchmark", 0))
        values = [cid, row.get("company_name")]
        for col in METRIC_COLUMNS:
            v = row.get(col)
            if pd.isna(v) if not isinstance(v, str) else False:
                v = None
            elif isinstance(v, float):
                v = round(v, 2)
            values.append(v)
        pct_row = pct_pivot.loc[cid] if cid in pct_pivot.index else pd.Series(dtype=float)
        pct_values = [pct_row.get(m) for m in metric_names]
        values.extend([round(p, 3) if pd.notna(p) else None for p in pct_values])

        ws.append(values)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.font = BODY_FONT
            if is_benchmark:
                cell.fill = GOLD

        pct_start_col = 2 + len(METRIC_COLUMNS) + 1
        for i, p in enumerate(pct_values):
            fill = _percentile_fill(p)
            if fill and not is_benchmark:
                ws.cell(row=row_idx, column=pct_start_col + i).fill = fill
        row_idx += 1

    # Summary row: peer-group median for each numeric metric column + percentile column.
    summary = ["MEDIAN", ""]
    for col in METRIC_COLUMNS:
        if col in numeric_metric_cols:
            med = pd.to_numeric(group_df[col], errors="coerce").median()
            summary.append(round(med, 2) if pd.notna(med) else None)
        else:
            summary.append(None)
    for m in metric_names:
        vals = pct_pivot[m] if m in pct_pivot.columns else pd.Series(dtype=float)
        med = vals.median()
        summary.append(round(med, 3) if pd.notna(med) else None)
    ws.append(summary)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = Font(name="Arial", bold=True, size=10)
        cell.fill = SUMMARY_FILL

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = max(10, len(headers[c - 1]) + 2)
    ws.freeze_panes = "C2"


def export_peer_comparison(snapshot_df: pd.DataFrame, percentiles_df: pd.DataFrame, out_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    assigned = snapshot_df[snapshot_df["peer_group_name"] != "No peer group assigned"]
    for group_name, grp in assigned.groupby("peer_group_name"):
        write_peer_group_sheet(wb, group_name, grp, percentiles_df)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
